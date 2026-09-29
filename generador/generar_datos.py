"""Genera los datos del caso RGM (ventas semanales de una cadena ficticia).

Salidas:
  datos/caso_rgm_datos.xlsx          -> lo que recibe el participante (data sucia)
  evaluador/datos_limpios.xlsx       -> la verdad sin errores (solo evaluador)
  evaluador/respuestas.json          -> cifras de referencia para la guía

Uso: python3 generador/generar_datos.py   (requiere openpyxl)
"""
import json
import math
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(2026)
ROOT = Path(__file__).resolve().parent.parent
WEEKS = range(1, 27)
START = date(2026, 1, 5)  # lunes
FORMATS = {"Hiper": 8, "Súper": 25, "Bodega": 60}  # número de tiendas


def monday(w):
    return START + timedelta(days=7 * (w - 1))


# sku, desc, categoría, marca, tipo, contenido, unidad, precio_reg, costo,
# unidades base por tienda por semana (H, S, B), elasticidad (H, S, B)
SKUS = [
    ("GAS-001", "Cola Tropical 2L", "Gaseosas", "Tropical", "Líder", 2, "L", 1450, 950, (110, 45, 22), (-2.0, -2.2, -2.8)),
    ("GAS-002", "Cola Tropical 3L", "Gaseosas", "Tropical", "Líder", 3, "L", 2350, 1500, (40, 15, 6), (-1.8, -1.8, -2.2)),
    ("GAS-003", "Precio Justo Cola 2L", "Gaseosas", "Precio Justo", "Marca Propia", 2, "L", 1050, 640, (35, 20, 18), (-1.5, -1.5, -1.8)),
    ("CAF-001", "Café Volcán 250g", "Café", "Volcán", "Líder", 250, "g", 2800, 1850, (45, 20, 9), (-1.0, -1.1, -1.4)),
    ("CAF-002", "Café Volcán 500g", "Café", "Volcán", "Líder", 500, "g", 5400, 3500, (18, 7, 2), (-1.0, -1.0, -1.2)),
    ("CAF-003", "Precio Justo Café 250g", "Café", "Precio Justo", "Marca Propia", 250, "g", 1950, 1250, (15, 12, 14), (-1.2, -1.2, -1.5)),
    ("ACE-001", "Aceite Girasol Dorado 1L", "Aceites", "Girasol Dorado", "Líder", 1, "L", 2200, 1500, (50, 22, 12), (-0.4, -0.5, -0.7)),
    ("ACE-002", "Aceite Palma Rica 1L", "Aceites", "Palma Rica", "Retador", 1, "L", 1850, 1300, (25, 14, 15), (-0.9, -0.9, -1.1)),
    ("ACE-003", "Precio Justo Aceite 1L", "Aceites", "Precio Justo", "Marca Propia", 1, "L", 1590, 1150, (15, 12, 18), (-1.0, -1.0, -1.2)),
    ("DET-001", "Detergente Blanco Max 1kg", "Detergentes", "Blanco Max", "Líder", 1, "kg", 2500, 1750, (55, 24, 12), (-1.5, -1.6, -2.0)),
    ("DET-002", "Detergente Espuma 1kg", "Detergentes", "Espuma", "Retador", 1, "kg", 2150, 1450, (30, 15, 12), (-1.3, -1.3, -1.6)),
    ("DET-003", "Precio Justo Detergente 1kg", "Detergentes", "Precio Justo", "Marca Propia", 1, "kg", 1690, 1150, (15, 14, 20), (-1.2, -1.2, -1.5)),
]
# SKU nuevo que aparece en ventas pero NO en el maestro (sin costo)
NEW_SKU = ("GAS-004", "Cola Tropical Zero 2L", "Gaseosas", "Tropical", "Líder", 2, "L", 1450, 960, (25, 8, 3), (-2.0, -2.2, -2.8))

GAS_PROMO_WEEKS = {13, 20}           # Cola Tropical 2L al 25 % (13 = Semana Santa)
DET_PROMO_WEEKS = {16, 17, 18, 19}   # Blanco Max al 30 %
CAF_UP_FROM, ACE_UP_FROM = 10, 8     # alzas de precio
STOCKOUT = ("ACE-003", "Súper", {18, 19, 20})


def price_for(sku, w, reg):
    if sku == "GAS-001" and w in GAS_PROMO_WEEKS:
        return 1090, "Sí", "Descuento 25%"
    if sku == "DET-001" and w in DET_PROMO_WEEKS:
        return 1750, "Sí", "Descuento 30%"
    if sku == "CAF-001" and w >= CAF_UP_FROM:
        return 3030, "No", ""
    if sku == "ACE-001" and w >= ACE_UP_FROM:
        return 2420, "No", ""
    return reg, "No", ""


def reg_price(sku, w, reg):
    # el precio regular de lista también sube con las alzas
    if sku == "CAF-001" and w >= CAF_UP_FROM:
        return 3030
    if sku == "ACE-001" and w >= ACE_UP_FROM:
        return 2420
    return reg


def multipliers(sku, fmt, w):
    m = 1.0
    cat = sku[:3]
    if cat == "GAS":  # estacionalidad Semana Santa (semana 13) y previa
        m *= {12: 1.10, 13: 1.35}.get(w, 1.0)
        if sku == "GAS-001" and w in {14, 21}:
            m *= 0.85  # despensa llena tras la promo
        if sku in {"GAS-002", "GAS-004"} and w in GAS_PROMO_WEEKS:
            m *= 0.80
        if sku == "GAS-003" and w in GAS_PROMO_WEEKS:
            m *= 0.70
    if sku == "DET-001" and w in {20, 21}:
        m *= {20: 0.80, 21: 0.92}[w]
    if sku == "DET-002" and w in DET_PROMO_WEEKS:
        m *= 0.75
    if sku == "DET-003" and w in DET_PROMO_WEEKS:
        m *= 0.80
    if sku in {"CAF-002", "CAF-003"} and w >= CAF_UP_FROM:
        m *= 1.06 if sku == "CAF-002" else 1.08
    if sku == "ACE-002" and fmt == "Súper" and w in STOCKOUT[2]:
        m *= 1.20
    return m


def generate_clean():
    rows = []
    for spec in SKUS + [NEW_SKU]:
        sku, desc, cat, marca, tipo, cont, uni, reg, cost, base, elas = spec
        for i, (fmt, n) in enumerate(FORMATS.items()):
            for w in WEEKS:
                if sku == "GAS-004" and w < 22:
                    continue
                if sku == STOCKOUT[0] and fmt == STOCKOUT[1] and w in STOCKOUT[2]:
                    continue  # quiebre: la fila simplemente no existe
                p, promo, tipo_promo = price_for(sku, w, reg)
                pr = reg_price(sku, w, reg)
                expected = base[i] * n * (p / reg) ** elas[i] * multipliers(sku, fmt, w)
                units = max(0, round(expected * math.exp(random.gauss(0, 0.06))))
                rows.append({
                    "Fecha": monday(w), "Semana": w, "Formato": fmt, "SKU": sku,
                    "Descripcion": desc, "Categoria": cat, "Precio_Regular": pr,
                    "Precio_Venta": p, "Promo": promo, "Tipo_Promo": tipo_promo,
                    "Unidades": units, "Venta": units * p, "Costo": cost,
                })
    return rows


def dirty(clean):
    """Inyecta errores y devuelve (filas_sucias, bitácora de errores)."""
    log = {}
    rows = [dict(r) for r in clean]
    idx = list(range(len(rows)))
    random.shuffle(idx)
    take = lambda k: [idx.pop() for _ in range(k)]

    # 1. unidades con un cero de más (Venta queda con el valor correcto)
    typo = take(3)
    for i in typo:
        rows[i]["Unidades"] = rows[i]["Unidades"] * 10
    log["unidades_x10"] = [(rows[i]["SKU"], rows[i]["Formato"], str(rows[i]["Fecha"])) for i in typo]

    # 2. precio de venta vacío (recuperable como Venta / Unidades)
    blank = take(12)
    for i in blank:
        rows[i]["Precio_Venta"] = None
    log["precio_vacio"] = len(blank)

    # 3. precio como texto con símbolo o coma decimal
    txt = take(25)
    for k, i in enumerate(txt):
        p = rows[i]["Precio_Venta"]
        rows[i]["Precio_Venta"] = f"₡{p:,}".replace(",", ".") if k % 2 else f"{p},00"
    log["precio_texto"] = len(txt)

    # 4. categoría, descripción y formato con variantes de escritura
    cat_var = {"Gaseosas": ["gaseosas", "GASEOSAS ", "Gaseosa"], "Café": ["Cafe", "CAFÉ", "café "],
               "Aceites": ["Aceite", "aceites", "ACEITES "], "Detergentes": ["Detergente", "detergentes", "DETERGENTES "]}
    fmt_var = {"Hiper": ["HIPER", "Hipermercado", "hiper "], "Súper": ["Super", "SUPER", "Supermercado"],
               "Bodega": ["BODEGA", "bodega ", "Bodega "]}
    for i in take(60):
        rows[i]["Categoria"] = random.choice(cat_var[rows[i]["Categoria"]])
    for i in take(45):
        rows[i]["Formato"] = random.choice(fmt_var[rows[i]["Formato"]])
    for i in take(40):
        d = rows[i]["Descripcion"]
        rows[i]["Descripcion"] = random.choice([d.upper(), d + "  ", " " + d.lower()])
    log["variantes_texto"] = "60 categoría, 45 formato, 40 descripción"

    # 5. fechas como texto
    for k, i in enumerate(take(40)):
        f = rows[i]["Fecha"]
        rows[i]["Fecha"] = f.strftime("%d/%m/%Y") if k % 2 else f.isoformat()
    log["fechas_texto"] = 40

    # 6. promo mal marcada: descuento real sin bandera y bandera sin descuento
    for r in rows:
        if r["SKU"] == "GAS-001" and r["Semana"] == 20 and r["Formato"].strip().lower().startswith("bod"):
            r["Promo"], r["Tipo_Promo"] = "No", ""
    fake = [i for i in take(4)]
    for i in fake:
        rows[i]["Promo"], rows[i]["Tipo_Promo"] = "Sí", "Descuento 10%"
    log["promo_mal_marcada"] = {"GAS-001 sem 20 Bodega sin bandera": 1,
                                "bandera Sí sin descuento": [(rows[i]["SKU"], str(rows[i]["Fecha"])) for i in fake]}

    # 7. devoluciones (unidades negativas) como filas extra
    returns = []
    for _ in range(6):
        base = clean[random.randrange(len(clean))]
        u = -random.randint(3, 12)
        returns.append({**base, "Unidades": u, "Venta": u * base["Precio_Venta"]})
    log["devoluciones"] = len(returns)

    # 8. duplicados: 14 exactos + 10 que solo se detectan tras estandarizar texto
    exact = [dict(rows[i]) for i in random.sample(range(len(rows)), 14)]
    near = []
    for i in random.sample(range(len(rows)), 10):
        r = dict(rows[i])
        r["Descripcion"] = r["Descripcion"].upper() + " "
        r["Formato"] = r["Formato"].upper()
        near.append(r)
    log["duplicados_exactos"], log["duplicados_tras_estandarizar"] = 14, 10

    out = rows + returns + exact + near
    random.shuffle(out)
    log["filas_entregadas"] = len(out)
    log["filas_limpias"] = len(clean)
    return out, log


COLS = ["Fecha", "Formato", "SKU", "Descripcion", "Categoria", "Precio_Regular",
        "Precio_Venta", "Promo", "Tipo_Promo", "Unidades", "Venta"]


def write_sheet(ws, header, rows):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16


def maestro_rows():
    return [[s[0], s[1], s[2], s[3], s[4], s[5], s[6], s[8]] for s in SKUS]


def write_participant(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    write_sheet(ws, COLS, [[r[c] for c in COLS] for r in rows])
    for cell in ws["A"][1:]:
        if isinstance(cell.value, date):
            cell.number_format = "yyyy-mm-dd"
    write_sheet(wb.create_sheet("Maestro_Productos"),
                ["SKU", "Descripcion", "Categoria", "Marca", "Tipo_Marca", "Contenido", "Unidad_Medida", "Costo_Unitario"],
                maestro_rows())
    write_sheet(wb.create_sheet("Tiendas"), ["Formato", "Numero_Tiendas"], [[k, v] for k, v in FORMATS.items()])
    write_sheet(wb.create_sheet("Diccionario"), ["Campo", "Descripción"], [
        ["Fecha", "Lunes de la semana de venta"],
        ["Formato", "Formato de tienda: Hiper, Súper o Bodega (descuento)"],
        ["SKU", "Código del producto"],
        ["Descripcion", "Nombre del producto"],
        ["Categoria", "Categoría comercial"],
        ["Precio_Regular", "Precio de lista en colones (₡), IVA incluido"],
        ["Precio_Venta", "Precio promedio cobrado en la semana (₡)"],
        ["Promo", "Sí / No: el producto tuvo actividad promocional"],
        ["Tipo_Promo", "Mecánica de la promoción"],
        ["Unidades", "Unidades vendidas en todas las tiendas del formato"],
        ["Venta", "Venta en colones (₡) = Unidades × Precio_Venta"],
        ["Costo_Unitario (Maestro)", "Costo de compra por unidad (₡)"],
        ["Numero_Tiendas (Tiendas)", "Tiendas activas por formato durante el periodo"],
    ])
    wb.save(path)


def write_clean(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas_Limpias"
    cols = COLS + ["Semana", "Costo"]
    write_sheet(ws, cols, [[r[c] for c in cols] for r in rows])
    wb.save(path)


# ---------- cifras de referencia para el evaluador ----------

def agg(rows, **filt):
    sel = [r for r in rows if all((r[k] in v) if isinstance(v, (set, list)) else r[k] == v for k, v in filt.items())]
    u = sum(r["Unidades"] for r in sel)
    v = sum(r["Venta"] for r in sel)
    m = sum(r["Venta"] - r["Unidades"] * r["Costo"] for r in sel)
    return u, v, m


def answers(clean):
    a = {}
    cats = {"GAS": "Gaseosas", "CAF": "Café", "ACE": "Aceites", "DET": "Detergentes"}
    a["categoria"] = {}
    for c in cats.values():
        u, v, m = agg(clean, Categoria=c)
        a["categoria"][c] = {"unidades": u, "venta": v, "margen": m, "margen_pct": round(100 * m / v, 1)}
    a["formato"] = {}
    for f, n in FORMATS.items():
        u, v, m = agg(clean, Formato=f)
        a["formato"][f] = {"venta": v, "venta_por_tienda_semana": round(v / n / 26), "margen_pct": round(100 * m / v, 1),
                           "share_marca_propia_pct": round(100 * agg([r for r in clean if r["SKU"] in {"GAS-003", "CAF-003", "ACE-003", "DET-003"}], Formato=f)[1] / v, 1)}

    def weekly(sku, weeks, fmt=None):
        f = {"SKU": sku, "Semana": set(weeks)}
        if fmt:
            f["Formato"] = fmt
        u, v, m = agg(clean, **f)
        return u / len(weeks), v / len(weeks), m / len(weeks)

    # Elasticidad GAS-001: semana 20 (limpia) vs 13 (Semana Santa) contra base 16-19
    base_w = [16, 17, 18, 19]
    ub = weekly("GAS-001", base_w)[0]
    pchg = 1090 / 1450 - 1
    e20 = (weekly("GAS-001", [20])[0] / ub - 1) / pchg
    e13 = (weekly("GAS-001", [13])[0] / ub - 1) / pchg
    # ajuste de estacionalidad con la marca propia no promocionada no sirve (fue canibalizada); usar GAS-003 semana 13 vs 20
    lg = lambda ur, pr: round(math.log(ur) / math.log(pr), 2)  # elasticidad log-log (punto)
    a["elasticidad_GAS001"] = {"semana20_vs_base": round(e20, 2), "semana13_vs_base_sin_ajustar": round(e13, 2),
                               "log_semana20": lg(weekly("GAS-001", [20])[0] / ub, 1090 / 1450),
                               "log_semana13": lg(weekly("GAS-001", [13])[0] / ub, 1090 / 1450),
                               "lift_unidades_sem20_pct": round(100 * (weekly("GAS-001", [20])[0] / ub - 1), 1),
                               "lift_unidades_sem13_pct": round(100 * (weekly("GAS-001", [13])[0] / ub - 1), 1)}
    a["elasticidad_GAS001_por_formato"] = {
        f: round((weekly("GAS-001", [20], f)[0] / weekly("GAS-001", base_w, f)[0] - 1) / pchg, 2) for f in FORMATS}

    # Promo detergente: 4 semanas vs base 10-15
    db = weekly("DET-001", range(10, 16))
    dp = weekly("DET-001", DET_PROMO_WEEKS)
    e_det = (dp[0] / db[0] - 1) / (1750 / 2500 - 1)
    promo_margin = dp[2] * 4 - db[2] * 4
    post = sum(weekly("DET-001", [w])[2] - db[2] for w in (20, 21))
    cani = 0
    for s in ("DET-002", "DET-003"):
        cani += (weekly(s, DET_PROMO_WEEKS)[2] - weekly(s, range(10, 16))[2]) * 4
    a["promo_DET001"] = {"elasticidad": round(e_det, 2), "elasticidad_log": round(math.log(dp[0] / db[0]) / math.log(0.7), 2), "lift_unidades_pct": round(100 * (dp[0] / db[0] - 1), 1),
                         "venta_incremental_4sem": round((dp[1] - db[1]) * 4),
                         "margen_incremental_DET001_4sem": round(promo_margin),
                         "margen_perdido_post_promo_sem20_21": round(post),
                         "margen_perdido_canibalizacion": round(cani),
                         "margen_neto_categoria": round(promo_margin + post + cani),
                         "margen_unitario_promo": 1750 - 1750}

    gb = weekly("GAS-001", base_w)
    gp = weekly("GAS-001", [20])
    g_post = weekly("GAS-001", [21])[2] - gb[2]
    g_cani = sum(weekly(s, [20])[2] - weekly(s, base_w)[2] for s in ("GAS-002", "GAS-003"))
    a["promo_GAS001_sem20"] = {"margen_incremental_GAS001": round(gp[2] - gb[2]),
                               "margen_perdido_post_promo_sem21": round(g_post),
                               "margen_perdido_canibalizacion": round(g_cani),
                               "margen_neto_categoria": round(gp[2] - gb[2] + g_post + g_cani),
                               "venta_incremental": round(gp[1] - gb[1])}

    for sku, up, newp, oldp in (("ACE-001", ACE_UP_FROM, 2420, 2200), ("CAF-001", CAF_UP_FROM, 3030, 2800)):
        pre, pos = weekly(sku, range(1, up)), weekly(sku, range(up, 27))
        a[f"alza_{sku}"] = {"precio_pct": round(100 * (newp / oldp - 1), 1),
                            "unidades_pct": round(100 * (pos[0] / pre[0] - 1), 1),
                            "venta_semanal_pct": round(100 * (pos[1] / pre[1] - 1), 1),
                            "margen_semanal_pct": round(100 * (pos[2] / pre[2] - 1), 1),
                            "elasticidad": round((pos[0] / pre[0] - 1) / (newp / oldp - 1), 2)}
    # migración dentro de café
    a["cafe_migracion"] = {s: round(100 * (weekly(s, range(10, 27))[0] / weekly(s, range(1, 10))[0] - 1), 1)
                           for s in ("CAF-002", "CAF-003")}

    pre = weekly("ACE-003", [w for w in range(12, 18)], "Súper")
    a["quiebre_ACE003_super"] = {"semanas": sorted(STOCKOUT[2]), "unidades_semana_base": round(pre[0]),
                                 "venta_perdida_estimada": round(pre[1] * 3), "margen_perdido_estimado": round(pre[2] * 3)}
    a["precio_por_litro"] = {"Cola Tropical 2L": 1450 / 2, "Cola Tropical 3L": round(2350 / 3, 1), "Precio Justo Cola 2L": 1050 / 2}
    a["GAS004"] = {"venta": agg(clean, SKU="GAS-004")[1], "unidades": agg(clean, SKU="GAS-004")[0]}
    return a


if __name__ == "__main__":
    clean = generate_clean()
    dirty_rows, log = dirty(clean)
    (ROOT / "datos").mkdir(exist_ok=True)
    (ROOT / "evaluador").mkdir(exist_ok=True)
    write_participant(dirty_rows, ROOT / "datos" / "caso_rgm_datos.xlsx")
    write_clean(clean, ROOT / "evaluador" / "datos_limpios.xlsx")
    res = {"errores_inyectados": log, **answers(clean)}
    (ROOT / "evaluador" / "respuestas.json").write_text(json.dumps(res, ensure_ascii=False, indent=2, default=str))
    # autocomprobación mínima
    assert log["filas_entregadas"] == log["filas_limpias"] + 6 + 14 + 10
    assert res["promo_DET001"]["margen_neto_categoria"] < 0, "la promo de detergente debe destruir margen"
    assert res["alza_ACE-001"]["margen_semanal_pct"] > 0, "el alza de aceite debe subir margen"
    assert res["elasticidad_GAS001"]["semana13_vs_base_sin_ajustar"] < res["elasticidad_GAS001"]["semana20_vs_base"]
    print(json.dumps(res, ensure_ascii=False, indent=2, default=str))
